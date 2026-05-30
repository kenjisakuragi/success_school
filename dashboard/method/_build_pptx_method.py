"""Marp Markdown → PPTX 変換スクリプト(桜木賢治セミナー専用)"""
import re
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

SRC = Path(__file__).parent / "method_v3_slides.md"
OUT = Path(__file__).parent / "sakuragi_method_v3_slides_v2.pptx"

# ===== Colors =====
PRIMARY = RGBColor(0xc9, 0x4f, 0x7c)
TITLE_DARK = RGBColor(0xb8, 0x45, 0x6e)
ACCENT = RGBColor(0xec, 0x99, 0xb3)
ACCENT_PALE = RGBColor(0xfc, 0xe4, 0xec)
BG_LIGHT = RGBColor(0xfd, 0xf5, 0xf8)
BG_WHITE = RGBColor(0xff, 0xff, 0xff)
BG_OFFER = RGBColor(0xff, 0xf8, 0xfb)
BG_SCIENCE = RGBColor(0xf5, 0xf3, 0xfa)
TEXT_DARK = RGBColor(0x2c, 0x2c, 0x2c)
TEXT_MUTED = RGBColor(0x77, 0x77, 0x77)
PURPLE = RGBColor(0x9c, 0x89, 0xc4)
HEADER_PINK = RGBColor(0xff, 0xf5, 0xf7)

FONT_JP = "Yu Gothic"
FONT_JP_BOLD = "Yu Gothic UI"

# ===== Slide size 16:9 (13.333 x 7.5 inch) =====
SW = Inches(13.333)
SH = Inches(7.5)

# ============= Parser =============

def parse_slides(md_path):
    text = md_path.read_text(encoding='utf-8')
    if text.startswith('---'):
        text = text.split('---', 2)[2]
    raw_slides = re.split(r'^---\s*$', text, flags=re.M)
    slides = []
    for raw in raw_slides:
        raw = raw.strip()
        if not raw:
            continue
        m = re.search(r'<!--\s*_class:\s*(\w+)\s*-->', raw)
        cls = m.group(1) if m else 'default'
        content = re.sub(r'<!--.*?-->', '', raw, flags=re.S).strip()
        slides.append({'class': cls, 'content': content})
    return slides

def parse_blocks(content):
    """Split content into blocks: heading, paragraph, bullets, table, blockquote."""
    lines = content.split('\n')
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        # Heading
        m = re.match(r'^(#{1,4})\s+(.+)$', line)
        if m:
            level = len(m.group(1))
            blocks.append({'type': 'heading', 'level': level, 'text': m.group(2).strip()})
            i += 1
            continue
        # Table
        if '|' in line and i+1 < len(lines) and re.match(r'^\s*\|?[\s\-:|]+\|?\s*$', lines[i+1]):
            tbl = []
            while i < len(lines) and '|' in lines[i]:
                row = lines[i].strip()
                if re.match(r'^\s*\|?[\s\-:|]+\|?\s*$', row):
                    i += 1
                    continue
                cells = [c.strip() for c in row.strip('|').split('|')]
                tbl.append(cells)
                i += 1
            blocks.append({'type': 'table', 'rows': tbl})
            continue
        # Bullet
        if re.match(r'^[\-\*]\s+', line) or re.match(r'^[🌸🌹🌺✅❌🔴⚪⚫▶︎]', line):
            items = []
            while i < len(lines):
                ln = lines[i].rstrip()
                m2 = re.match(r'^[\-\*]\s+(.+)$', ln)
                if m2:
                    items.append(m2.group(1).strip())
                    i += 1
                elif ln and re.match(r'^[🌸🌹🌺✅❌🔴⚪⚫▶︎]', ln):
                    items.append(ln.strip())
                    i += 1
                elif not ln:
                    i += 1
                    if i < len(lines) and not (re.match(r'^[\-\*]\s+', lines[i]) or re.match(r'^[🌸🌹🌺✅❌🔴⚪⚫▶︎]', lines[i])):
                        break
                else:
                    break
            blocks.append({'type': 'bullets', 'items': items})
            continue
        # Blockquote
        if line.startswith('>'):
            quote_lines = []
            while i < len(lines) and lines[i].startswith('>'):
                quote_lines.append(lines[i][1:].strip())
                i += 1
            blocks.append({'type': 'quote', 'text': '\n'.join(quote_lines).strip()})
            continue
        # Paragraph
        para_lines = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r'^[#\-\*>]', lines[i]) and '|' not in lines[i]:
            para_lines.append(lines[i].rstrip())
            i += 1
        blocks.append({'type': 'paragraph', 'text': '\n'.join(para_lines).strip()})
    return blocks

# ============= Helpers =============

def set_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_rect(slide, left, top, width, height, fill_color, line_color=None, line_width=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill_color
    if line_color:
        shp.line.color.rgb = line_color
        if line_width:
            shp.line.width = line_width
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp

def add_text(slide, left, top, width, height, text, *, size=24, color=TEXT_DARK,
             bold=False, italic=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             font=FONT_JP, line_spacing=1.2):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    lines = text.split('\n') if isinstance(text, str) else text
    for idx, ln in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = ln
        run.font.name = font
        # Also set East Asian font
        rPr = run._r.get_or_add_rPr()
        ea = rPr.find(qn('a:ea'))
        if ea is None:
            ea = etree.SubElement(rPr, qn('a:ea'))
        ea.set('typeface', font)
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.italic = italic
    return tb

def add_header_strip(slide, color=HEADER_PINK, height=Inches(0.18)):
    add_rect(slide, 0, 0, SW, height, color)

def add_footer(slide):
    add_text(slide, Inches(0.5), Inches(7.05), Inches(8), Inches(0.35),
             "© 2026 桜木賢治チャンネル プロジェクト",
             size=10, color=TEXT_MUTED, align=PP_ALIGN.LEFT)

def add_page_num(slide, num, total):
    add_text(slide, Inches(12.3), Inches(7.05), Inches(0.8), Inches(0.35),
             f"{num} / {total}",
             size=10, color=TEXT_MUTED, align=PP_ALIGN.RIGHT)

def add_title_bar(slide, title, color=PRIMARY):
    # Left vertical accent bar + title text
    add_rect(slide, Inches(0.5), Inches(0.6), Inches(0.1), Inches(0.65), color)
    add_text(slide, Inches(0.75), Inches(0.5), Inches(12), Inches(0.85),
             title, size=32, color=TITLE_DARK, bold=True,
             font=FONT_JP_BOLD, anchor=MSO_ANCHOR.MIDDLE)

# Inline markdown to plain text (strip **bold** markers but mark them)
def parse_inline(text):
    """Return list of (text, bold) tuples handling **bold** markers."""
    parts = []
    i = 0
    while i < len(text):
        m = re.search(r'\*\*(.+?)\*\*', text[i:])
        if not m:
            parts.append((text[i:], False))
            break
        if m.start() > 0:
            parts.append((text[i:i+m.start()], False))
        parts.append((m.group(1), True))
        i += m.end()
    return [p for p in parts if p[0]]

def add_rich_text(slide, left, top, width, height, content_lines, *,
                  size=22, color=TEXT_DARK, align=PP_ALIGN.LEFT,
                  anchor=MSO_ANCHOR.TOP, font=FONT_JP, line_spacing=1.3,
                  bullet_color=None):
    """content_lines: list of strings, each may have **bold** markers."""
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = anchor
    for idx, line in enumerate(content_lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        parts = parse_inline(line) if line else [('', False)]
        for ptext, pbold in parts:
            run = p.add_run()
            run.text = ptext
            run.font.name = font
            rPr = run._r.get_or_add_rPr()
            ea = rPr.find(qn('a:ea'))
            if ea is None:
                ea = etree.SubElement(rPr, qn('a:ea'))
            ea.set('typeface', font)
            run.font.size = Pt(size)
            if pbold:
                run.font.bold = True
                run.font.color.rgb = PRIMARY
            else:
                run.font.color.rgb = color
    return tb

# ============= Slide renderers =============

def render_default(slide, content, header_color=HEADER_PINK):
    set_bg(slide, BG_WHITE)
    add_header_strip(slide, header_color)
    blocks = parse_blocks(content)
    # Find first H1/H2 as title
    title = None
    body_blocks = []
    for b in blocks:
        if b['type'] == 'heading' and b['level'] <= 2 and title is None:
            title = b['text']
        else:
            body_blocks.append(b)
    if title:
        add_title_bar(slide, title)
    render_body(slide, body_blocks, top=Inches(1.5))

def render_title(slide, content):
    # Gradient-feel: two rects + accent
    set_bg(slide, BG_LIGHT)
    add_rect(slide, 0, 0, SW, Inches(7.5), BG_LIGHT)
    # Decorative top arc band
    add_rect(slide, 0, 0, SW, Inches(0.4), ACCENT)
    add_rect(slide, 0, Inches(7.1), SW, Inches(0.4), ACCENT)
    blocks = parse_blocks(content)
    # Collect headings & paragraphs
    titles = [b['text'] for b in blocks if b['type'] == 'heading' and b['level'] == 1]
    subs = [b['text'] for b in blocks if b['type'] == 'heading' and b['level'] == 2]
    paras = [b['text'] for b in blocks if b['type'] == 'paragraph']
    quotes = [b['text'] for b in blocks if b['type'] == 'quote']
    y = Inches(1.2)
    for t in titles:
        add_text(slide, Inches(0.5), y, Inches(12.3), Inches(1.2),
                 t, size=56, color=PRIMARY, bold=True,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                 font=FONT_JP_BOLD)
        y += Inches(1.2)
    for s in subs:
        add_text(slide, Inches(0.5), y, Inches(12.3), Inches(0.9),
                 s, size=32, color=TITLE_DARK, bold=True,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                 font=FONT_JP_BOLD)
        y += Inches(0.9)
    for p in paras + quotes:
        clean = re.sub(r'<.*?>', '', p)
        add_text(slide, Inches(0.5), y, Inches(12.3), Inches(0.7),
                 clean, size=22, color=TEXT_DARK,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        y += Inches(0.7)

def render_part(slide, content):
    # Section divider
    set_bg(slide, ACCENT_PALE)
    add_rect(slide, 0, 0, SW, Inches(7.5), ACCENT_PALE)
    # Big pink stripe
    add_rect(slide, 0, Inches(2.8), SW, Inches(1.9), ACCENT)
    blocks = parse_blocks(content)
    titles = [b['text'] for b in blocks if b['type'] == 'heading' and b['level'] == 1]
    subs = [b['text'] for b in blocks if b['type'] == 'heading' and b['level'] == 2]
    paras = [b['text'] for b in blocks if b['type'] == 'paragraph']
    if titles:
        add_text(slide, Inches(0.5), Inches(2.95), Inches(12.3), Inches(0.9),
                 titles[0], size=56, color=BG_WHITE, bold=True,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                 font=FONT_JP_BOLD)
    if subs:
        add_text(slide, Inches(0.5), Inches(3.85), Inches(12.3), Inches(0.8),
                 subs[0], size=30, color=BG_WHITE, bold=False,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                 font=FONT_JP_BOLD)
    if paras:
        add_text(slide, Inches(0.5), Inches(5.0), Inches(12.3), Inches(0.6),
                 paras[0], size=22, color=TITLE_DARK,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

def render_work(slide, content):
    set_bg(slide, BG_LIGHT)
    # Thick left border
    add_rect(slide, 0, 0, Inches(0.35), Inches(7.5), ACCENT)
    add_header_strip(slide, HEADER_PINK)
    blocks = parse_blocks(content)
    title = None
    body_blocks = []
    for b in blocks:
        if b['type'] == 'heading' and b['level'] <= 2 and title is None:
            title = b['text']
        else:
            body_blocks.append(b)
    if title:
        add_rect(slide, Inches(0.65), Inches(0.6), Inches(0.1), Inches(0.65), PRIMARY)
        add_text(slide, Inches(0.9), Inches(0.5), Inches(12), Inches(0.85),
                 title, size=32, color=TITLE_DARK, bold=True,
                 font=FONT_JP_BOLD, anchor=MSO_ANCHOR.MIDDLE)
    render_body(slide, body_blocks, top=Inches(1.5), left=Inches(0.75))

def render_offer(slide, content):
    set_bg(slide, BG_OFFER)
    # Border frame
    border = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
                                     Inches(0.2), Inches(0.2),
                                     Inches(12.933), Inches(7.1))
    border.fill.background()
    border.line.color.rgb = ACCENT
    border.line.width = Pt(3)
    border.shadow.inherit = False
    blocks = parse_blocks(content)
    title = None
    body_blocks = []
    for b in blocks:
        if b['type'] == 'heading' and b['level'] <= 2 and title is None:
            title = b['text']
        else:
            body_blocks.append(b)
    if title:
        add_rect(slide, Inches(0.7), Inches(0.65), Inches(0.1), Inches(0.7), PRIMARY)
        add_text(slide, Inches(0.95), Inches(0.55), Inches(12), Inches(0.9),
                 title, size=34, color=PRIMARY, bold=True,
                 font=FONT_JP_BOLD, anchor=MSO_ANCHOR.MIDDLE)
    render_body(slide, body_blocks, top=Inches(1.6), left=Inches(0.8), width=Inches(11.7))

def render_science(slide, content):
    set_bg(slide, BG_SCIENCE)
    add_rect(slide, 0, 0, Inches(0.25), Inches(7.5), PURPLE)
    add_header_strip(slide, HEADER_PINK)
    blocks = parse_blocks(content)
    title = None
    body_blocks = []
    for b in blocks:
        if b['type'] == 'heading' and b['level'] <= 2 and title is None:
            title = b['text']
        else:
            body_blocks.append(b)
    if title:
        add_rect(slide, Inches(0.6), Inches(0.6), Inches(0.1), Inches(0.65), PURPLE)
        add_text(slide, Inches(0.85), Inches(0.5), Inches(12), Inches(0.85),
                 title, size=32, color=PURPLE, bold=True,
                 font=FONT_JP_BOLD, anchor=MSO_ANCHOR.MIDDLE)
    render_body(slide, body_blocks, top=Inches(1.5), left=Inches(0.7))

# ============= Body block renderer =============

def render_body(slide, blocks, *, top=Inches(1.5), left=Inches(0.7), width=Inches(11.9)):
    y = top
    max_y = Inches(6.9)
    for b in blocks:
        if y >= max_y:
            break
        if b['type'] == 'heading':
            sz = {2: 30, 3: 26, 4: 22}.get(b['level'], 22)
            h = Inches(0.6) if b['level'] >= 3 else Inches(0.7)
            add_text(slide, left, y, width, h, b['text'],
                     size=sz, color=TITLE_DARK, bold=True,
                     font=FONT_JP_BOLD, anchor=MSO_ANCHOR.MIDDLE)
            y += h + Inches(0.05)
        elif b['type'] == 'paragraph':
            # Detect line count for sizing
            lines = [l for l in b['text'].split('\n') if l.strip()]
            # Strip span tags
            cleaned = []
            for ln in lines:
                ln = re.sub(r'<.*?>', '', ln)
                cleaned.append(ln)
            n = len(cleaned)
            est_h = Inches(0.55 * max(n, 1) + 0.1)
            est_h = min(est_h, max_y - y)
            add_rich_text(slide, left, y, width, est_h, cleaned,
                          size=24, color=TEXT_DARK, line_spacing=1.35)
            y += est_h + Inches(0.1)
        elif b['type'] == 'bullets':
            items = b['items']
            n = len(items)
            # Estimated height
            est_h = Inches(0.55 * n + 0.2)
            est_h = min(est_h, max_y - y)
            # Render as rich text with bullet markers preserved (emoji bullets stay)
            display = []
            for it in items:
                # If item doesn't start with emoji, prefix with • dot
                if not re.match(r'^[🌸🌹🌺✅❌🔴⚪⚫▶︎]', it):
                    display.append(f"・ {it}")
                else:
                    display.append(it)
            add_rich_text(slide, left, y, width, est_h, display,
                          size=24, color=TEXT_DARK, line_spacing=1.45)
            y += est_h + Inches(0.1)
        elif b['type'] == 'table':
            rows = b['rows']
            if not rows:
                continue
            ncols = max(len(r) for r in rows)
            nrows = len(rows)
            # auto height
            row_h = Inches(0.5)
            tbl_h = row_h * nrows
            tbl_w = width
            if y + tbl_h > max_y:
                tbl_h = max_y - y
            shape = slide.shapes.add_table(nrows, ncols, left, y, tbl_w, tbl_h)
            tbl = shape.table
            for ri, row in enumerate(rows):
                for ci in range(ncols):
                    cell = tbl.cell(ri, ci)
                    txt = row[ci] if ci < len(row) else ''
                    txt = re.sub(r'\*\*(.+?)\*\*', r'\1', txt)
                    cell.text = ''
                    tf = cell.text_frame
                    tf.margin_left = Inches(0.08)
                    tf.margin_right = Inches(0.08)
                    tf.margin_top = Inches(0.04)
                    tf.margin_bottom = Inches(0.04)
                    p = tf.paragraphs[0]
                    p.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER
                    run = p.add_run()
                    run.text = txt
                    run.font.name = FONT_JP
                    rPr = run._r.get_or_add_rPr()
                    ea = rPr.find(qn('a:ea'))
                    if ea is None:
                        ea = etree.SubElement(rPr, qn('a:ea'))
                    ea.set('typeface', FONT_JP)
                    is_bold_cell = ('**' in (row[ci] if ci < len(row) else ''))
                    if ri == 0:
                        run.font.size = Pt(18)
                        run.font.bold = True
                        run.font.color.rgb = TITLE_DARK
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = ACCENT_PALE
                    else:
                        run.font.size = Pt(17)
                        run.font.color.rgb = TEXT_DARK
                        if is_bold_cell:
                            run.font.bold = True
                            run.font.color.rgb = PRIMARY
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = BG_WHITE if ri % 2 == 1 else BG_LIGHT
            y += tbl_h + Inches(0.15)
        elif b['type'] == 'quote':
            txt = b['text']
            est_h = Inches(0.55 * max(len(txt.split('\n')), 1) + 0.3)
            est_h = min(est_h, max_y - y)
            # Background tint
            add_rect(slide, left, y, width, est_h, ACCENT_PALE)
            add_rect(slide, left, y, Inches(0.08), est_h, ACCENT)
            add_text(slide, left + Inches(0.2), y + Inches(0.05),
                     width - Inches(0.3), est_h - Inches(0.1),
                     txt, size=22, color=TITLE_DARK, italic=False,
                     anchor=MSO_ANCHOR.MIDDLE)
            y += est_h + Inches(0.15)

# ============= Main =============

def build():
    slides = parse_slides(SRC)
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    blank_layout = prs.slide_layouts[6]
    total = len(slides)
    for i, sl in enumerate(slides, 1):
        s = prs.slides.add_slide(blank_layout)
        cls = sl['class']
        if cls == 'title':
            render_title(s, sl['content'])
        elif cls == 'part':
            render_part(s, sl['content'])
        elif cls == 'work':
            render_work(s, sl['content'])
            add_page_num(s, i, total)
        elif cls == 'offer':
            render_offer(s, sl['content'])
            add_page_num(s, i, total)
        elif cls == 'science':
            render_science(s, sl['content'])
            add_page_num(s, i, total)
        else:
            render_default(s, sl['content'])
            add_page_num(s, i, total)
    prs.save(OUT)
    print(f"Generated {total} slides -> {OUT}")

if __name__ == '__main__':
    build()
