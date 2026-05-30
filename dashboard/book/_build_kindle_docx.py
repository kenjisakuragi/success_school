"""
Kindle入稿用 .docx 生成スクリプト

桜木賢治『科学的に正しい引き寄せ』第1部
全14ファイル(プロローグ〜エピローグ+巻末)をひとつの .docx にまとめる。

KDP(Amazon Kindle Direct Publishing)は .docx を直接受け付け、
Kindle Previewerで自動的にリフロー型電子書籍に変換される。

使い方:
    python _build_kindle_docx.py
"""
from __future__ import annotations

import re
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = Path(__file__).resolve().parent

# 章ファイル順序
CHAPTER_FILES = [
    ("part1_ch00_prologue.md", "プロローグ"),
    ("part1_ch01_origin.md", "第1章"),
    ("part1_ch02_half_truth.md", "第2章"),
    ("part1_ch03_belief_brain.md", "第3章"),
    ("part1_ch04_writing.md", "第4章"),
    ("part1_ch05_body.md", "第5章"),
    ("part1_ch06_woop.md", "第6章"),
    ("part1_ch07_three_pillars.md", "第7章"),
    ("part1_ch08_manifest_ai.md", "第8章"),
    ("part1_ch09_five_actions.md", "第9章"),
    ("part1_ch10_day1_10.md", "第10章"),
    ("part1_ch11_day11_20.md", "第11章"),
    ("part1_ch12_day21_30.md", "第12章"),
    ("part1_ch99_epilogue_and_appendix.md", "エピローグ+巻末"),
]

OUT_PATH = HERE / "SAKURAGI_BOOK_PART1_DRAFT1.docx"

JP_FONT = "Yu Mincho"  # 紙書籍想定の明朝体


def add_run_with_formatting(paragraph, text: str, bold=False, italic=False):
    """段落に書式付きのrunを追加"""
    run = paragraph.add_run(text)
    run.font.name = JP_FONT
    # 日本語フォント明示指定(Word対応)
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), JP_FONT)
    run.bold = bold
    run.italic = italic
    return run


def parse_inline_markdown(paragraph, text: str):
    """**bold** と *italic* を含むテキストを段落に追加"""
    # Match **bold** then leftover
    pattern = re.compile(r'(\*\*([^*]+)\*\*|\*([^*\s][^*]*?)\*|`([^`]+)`)')
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            add_run_with_formatting(paragraph, text[pos:m.start()])
        if m.group(2) is not None:  # **bold**
            add_run_with_formatting(paragraph, m.group(2), bold=True)
        elif m.group(3) is not None:  # *italic*
            add_run_with_formatting(paragraph, m.group(3), italic=True)
        elif m.group(4) is not None:  # `code`
            add_run_with_formatting(paragraph, m.group(4))  # plain
        pos = m.end()
    if pos < len(text):
        add_run_with_formatting(paragraph, text[pos:])


def add_pagebreak(doc):
    p = doc.add_paragraph()
    r = p.add_run()
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    r._element.append(br)


def process_md_lines(doc, lines: list[str]):
    """Markdown行リストを順次docxへ書き込む"""
    i = 0
    in_code_block = False
    code_lines: list[str] = []

    while i < len(lines):
        line = lines[i].rstrip()

        # ----- Code block (```) -----
        if line.strip().startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_lines = []
            else:
                # close
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.5)
                add_run_with_formatting(p, "\n".join(code_lines))
                in_code_block = False
            i += 1
            continue
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # ----- Skip empty -----
        if not line.strip():
            i += 1
            continue

        # ----- Front matter / metadata comments / horizontal rules -----
        if line.strip() == "---":
            i += 1
            continue

        # ----- Headings -----
        m = re.match(r'^(#{1,4})\s+(.+)$', line)
        if m:
            level = len(m.group(1))
            heading_text = m.group(2).strip()
            # Strip leading emojis if any (📕, 🌸, etc.)
            heading_text = re.sub(r'^[📕📚📖🌸🌹🎯🤖🧠✍️⚡👑🔁🎤🆓📦⏰🥇🥈🥉⚪🌟💰🎉]\s*', '', heading_text)
            doc.add_heading(heading_text, level=min(level, 4))
            i += 1
            continue

        # ----- Blockquote -----
        if line.startswith(">"):
            q_lines = []
            while i < len(lines) and lines[i].rstrip().startswith(">"):
                q_lines.append(lines[i].rstrip().lstrip(">").strip())
                i += 1
            quote_text = " ".join(q_lines).strip()
            p = doc.add_paragraph(style="Intense Quote")
            parse_inline_markdown(p, quote_text)
            continue

        # ----- Tables (markdown pipe tables) -----
        if line.startswith("|") and (i + 1 < len(lines)) and re.match(r'^\s*\|?[\s\-:|]+\|?\s*$', lines[i+1]):
            # gather table rows
            table_rows: list[list[str]] = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                row = lines[i].strip()
                if re.match(r'^\s*\|?[\s\-:|]+\|?\s*$', row):
                    i += 1
                    continue
                cells = [c.strip() for c in row.strip('|').split('|')]
                table_rows.append(cells)
                i += 1
            if table_rows:
                ncols = max(len(r) for r in table_rows)
                tbl = doc.add_table(rows=len(table_rows), cols=ncols)
                tbl.style = 'Light Grid Accent 1'
                for ri, row in enumerate(table_rows):
                    for ci in range(ncols):
                        cell = tbl.rows[ri].cells[ci]
                        cell_para = cell.paragraphs[0]
                        text = row[ci] if ci < len(row) else ""
                        parse_inline_markdown(cell_para, text)
                doc.add_paragraph()  # spacing
            continue

        # ----- List items -----
        m = re.match(r'^(\s*)([-*+])\s+(.+)$', line)
        if m:
            indent = len(m.group(1)) // 2
            text = m.group(3)
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Cm(0.5 + indent * 0.5)
            parse_inline_markdown(p, text)
            i += 1
            continue
        m = re.match(r'^(\s*)\d+\.\s+(.+)$', line)
        if m:
            text = m.group(2)
            p = doc.add_paragraph(style="List Number")
            parse_inline_markdown(p, text)
            i += 1
            continue

        # ----- Regular paragraph (combine continuous lines until blank) -----
        para_lines = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i].rstrip()
            if not nxt or re.match(r'^(#{1,4}\s|\||>|```|\s*[-*+]\s|---$)', nxt) or re.match(r'^\d+\.\s', nxt):
                break
            para_lines.append(nxt)
            i += 1
        para_text = "".join(para_lines)
        p = doc.add_paragraph()
        parse_inline_markdown(p, para_text)


def main():
    doc = Document()

    # Default style
    style = doc.styles['Normal']
    style.font.name = JP_FONT
    style.font.size = Pt(10.5)
    rPr = style.element.rPr
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), JP_FONT)

    # Title page
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("科学的に正しい引き寄せ")
    title_run.font.name = JP_FONT
    title_run.font.size = Pt(28)
    title_run.bold = True

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub.add_run("― マーフィーが届けたかった、本当の引き寄せ ―")
    sub_run.font.name = JP_FONT
    sub_run.font.size = Pt(14)

    doc.add_paragraph()
    doc.add_paragraph()

    author = doc.add_paragraph()
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author_run = author.add_run("桜木 賢治")
    author_run.font.name = JP_FONT
    author_run.font.size = Pt(16)

    org = doc.add_paragraph()
    org.alignment = WD_ALIGN_PARAGRAPH.CENTER
    org_run = org.add_run("日本マニフェスティング協会 代表")
    org_run.font.name = JP_FONT
    org_run.font.size = Pt(11)

    add_pagebreak(doc)

    # 各章を読み込み
    total_chars = 0
    for fname, label in CHAPTER_FILES:
        fp = HERE / fname
        if not fp.exists():
            print(f"[skip] {fname} not found", file=sys.stderr)
            continue
        text = fp.read_text(encoding='utf-8')
        # Skip metadata block at top (first H1 might be like "# 📕 第X章...")
        # Drop the very top metadata lines before the first heading
        lines = text.split("\n")
        # find first H1
        start = 0
        for k, ln in enumerate(lines):
            if ln.startswith("# "):
                start = k
                break
        chapter_lines = lines[start:]
        # filter out metadata/file-level lines: **第1稿:** / **文体方針:**
        chapter_lines = [
            ln for ln in chapter_lines
            if not ln.startswith("**第1稿:**")
            and not ln.startswith("**文体方針:**")
            and not ln.startswith("*プロローグ・")
            and not ln.startswith("*第")
            and not ln.startswith("*エピローグ・巻末資料")
        ]
        process_md_lines(doc, chapter_lines)
        add_pagebreak(doc)
        total_chars += sum(len(ln) for ln in chapter_lines)
        print(f"  OK: {fname} ({label}) - {sum(len(ln) for ln in chapter_lines):,} chars")

    doc.save(OUT_PATH)
    print(f"\n✨ 完成: {OUT_PATH}")
    print(f"   累計文字数(概算): {total_chars:,}")


if __name__ == "__main__":
    main()
